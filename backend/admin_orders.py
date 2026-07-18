"""
Admin — order book (pending + completed) + bulk actions + PDF/XLSX exports.
Order lifecycle:
  CONFIRMED (paid) -> SHIPPED -> DELIVERED
Pending = status in {"CONFIRMED"} (paid but not yet shipped).
Completed = filtered by date range (all statuses).
"""
import io
import json
from datetime import datetime, timezone, timedelta
from typing import Optional, List
from pathlib import Path

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from pydantic import BaseModel

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

from admin_auth import require_admin
from admin_pdf import (
    build_addresses_pdf,
    build_pending_orders_pdf,
    build_invoice_pdf,
)


router = APIRouter(prefix="/admin/orders", tags=["admin-orders"])


ALLOWED_SORT = {"created_at", "city", "totalQuantity", "total", "orderNumber"}


class MarkShippedRequest(BaseModel):
    orderNumbers: List[str]


class PrintAddressesRequest(BaseModel):
    orderNumbers: List[str]


def _sort_orders(orders: list[dict], sort: str, direction: str) -> list[dict]:
    reverse = direction.lower() == "desc"
    if sort == "city":
        return sorted(orders, key=lambda o: (o.get("address", {}).get("city") or "").lower(), reverse=reverse)
    if sort == "totalQuantity":
        return sorted(orders, key=lambda o: sum(i.get("quantity", 0) for i in o.get("items", [])), reverse=reverse)
    if sort == "total":
        return sorted(orders, key=lambda o: float(o.get("total", 0)), reverse=reverse)
    if sort == "orderNumber":
        return sorted(orders, key=lambda o: o.get("orderNumber", ""), reverse=reverse)
    # default created_at
    return sorted(orders, key=lambda o: o.get("created_at", ""), reverse=reverse)


def register(api: APIRouter, db):

    async def _completed_orders(from_date, to_date):
        # Default: last 30 days
        now = datetime.now(timezone.utc)
        to_dt = datetime.fromisoformat(to_date) if to_date else now
        from_dt = datetime.fromisoformat(from_date) if from_date else (now - timedelta(days=30))

        cursor = db.orders.find(
            {"created_at": {"$gte": from_dt.isoformat(), "$lte": to_dt.isoformat()}},
            {"_id": 0},
        )
        orders = [o async for o in cursor]
        orders.sort(key=lambda o: o.get("created_at", ""), reverse=True)
        return {
            "count": len(orders),
            "orders": orders,
            "from": from_dt.isoformat(),
            "to": to_dt.isoformat(),
        }

    @router.get("/pending")
    async def list_pending(
        sort: str = Query("created_at"),
        direction: str = Query("desc"),
        _: dict = Depends(require_admin),
    ):
        if sort not in ALLOWED_SORT:
            raise HTTPException(status_code=400, detail=f"Invalid sort field. Allowed: {sorted(ALLOWED_SORT)}")
        cursor = db.orders.find({"status": "CONFIRMED"}, {"_id": 0})
        orders = [o async for o in cursor]
        orders = _sort_orders(orders, sort, direction)
        return {"count": len(orders), "orders": orders, "sort": sort, "direction": direction}

    @router.get("/completed")
    async def list_completed(
        from_date: Optional[str] = Query(None),
        to_date: Optional[str] = Query(None),
        _: dict = Depends(require_admin),
    ):
        return await _completed_orders(from_date, to_date)

    @router.post("/mark-shipped")
    async def mark_shipped(req: MarkShippedRequest, admin=Depends(require_admin)):
        if not req.orderNumbers:
            raise HTTPException(status_code=400, detail="Provide at least one order number")
        now = datetime.now(timezone.utc).isoformat()
        result = await db.orders.update_many(
            {"orderNumber": {"$in": req.orderNumbers}, "status": "CONFIRMED"},
            {"$set": {"status": "SHIPPED", "shipped_at": now, "shipped_by": admin.get("sub")}},
        )
        return {"success": True, "shipped": result.modified_count, "requested": len(req.orderNumbers)}

    @router.post("/print-address")
    async def print_addresses(req: PrintAddressesRequest, _: dict = Depends(require_admin)):
        if not req.orderNumbers:
            raise HTTPException(status_code=400, detail="Provide at least one order number")
        cursor = db.orders.find({"orderNumber": {"$in": req.orderNumbers}}, {"_id": 0})
        orders = [o async for o in cursor]
        if not orders:
            raise HTTPException(status_code=404, detail="No matching orders")
        # Preserve requested order
        by_num = {o["orderNumber"]: o for o in orders}
        ordered = [by_num[n] for n in req.orderNumbers if n in by_num]
        pdf = build_addresses_pdf(ordered)
        return StreamingResponse(
            io.BytesIO(pdf),
            media_type="application/pdf",
            headers={"Content-Disposition": f'attachment; filename="addresses_{len(ordered)}.pdf"'},
        )

    @router.get("/print-pending-pdf")
    async def print_pending_pdf(
        sort: str = Query("created_at"),
        direction: str = Query("desc"),
        _: dict = Depends(require_admin),
    ):
        cursor = db.orders.find({"status": "CONFIRMED"}, {"_id": 0})
        orders = [o async for o in cursor]
        orders = _sort_orders(orders, sort, direction)
        pdf = build_pending_orders_pdf(orders)
        return StreamingResponse(
            io.BytesIO(pdf),
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="pending_orders.pdf"'},
        )

    @router.post("/{order_number}/send-invoice")
    async def send_invoice(order_number: str, _: dict = Depends(require_admin)):
        order = await db.orders.find_one({"orderNumber": order_number}, {"_id": 0})
        if not order:
            raise HTTPException(status_code=404, detail="Order not found")
        # Build PDF and store it. Real delivery via WhatsApp/SMS wires up at Twilio go-live.
        pdf = build_invoice_pdf(order)
        invoice_dir = Path(__file__).parent / "uploads" / "invoices"
        invoice_dir.mkdir(parents=True, exist_ok=True)
        target = invoice_dir / f"{order_number}.pdf"
        with open(target, "wb") as f:
            f.write(pdf)
        # Log intent-to-send (mocked)
        record = {
            "orderNumber": order_number,
            "phone": order.get("phone"),
            "invoicePath": str(target),
            "channel": "WHATSAPP",
            "provider": "MOCK",
            "status": "QUEUED",
            "queued_at": datetime.now(timezone.utc).isoformat(),
        }
        await db.invoice_dispatch_log.insert_one(dict(record))
        return {"success": True, "queued": True, "invoiceUrl": f"/api/admin/orders/{order_number}/invoice.pdf"}

    @router.get("/{order_number}/invoice.pdf")
    async def get_invoice_pdf(order_number: str, _: dict = Depends(require_admin)):
        order = await db.orders.find_one({"orderNumber": order_number}, {"_id": 0})
        if not order:
            raise HTTPException(status_code=404, detail="Order not found")
        pdf = build_invoice_pdf(order)
        return StreamingResponse(
            io.BytesIO(pdf),
            media_type="application/pdf",
            headers={"Content-Disposition": f'inline; filename="invoice_{order_number}.pdf"'},
        )

    @router.get("/completed/export.xlsx")
    async def export_completed_xlsx(
        from_date: Optional[str] = Query(None),
        to_date: Optional[str] = Query(None),
        _: dict = Depends(require_admin),
    ):
        orders = (await _completed_orders(from_date, to_date))["orders"]

        wb = Workbook()
        ws = wb.active
        ws.title = "Completed Orders"
        header = ["Order #", "Date", "Status", "Customer", "Phone", "City", "State",
                  "Items (name x qty)", "Subtotal", "Shipping", "Total"]
        ws.append(header)
        header_fill = PatternFill("solid", fgColor="174C2B")
        border = Border(left=Side(style="thin", color="EADFC7"),
                        right=Side(style="thin", color="EADFC7"),
                        top=Side(style="thin", color="EADFC7"),
                        bottom=Side(style="thin", color="EADFC7"))
        for cell in ws[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = border

        for o in orders:
            addr = o.get("address", {})
            items = ", ".join(
                f"{i.get('name','')} {i.get('variant','')} x {i.get('quantity',1)}"
                for i in o.get("items", [])
            )
            ws.append([
                o.get("orderNumber", ""),
                (o.get("created_at", "") or "")[:10],
                o.get("status", ""),
                addr.get("name", ""),
                f"+91 {o.get('phone', '')}",
                addr.get("city", ""),
                addr.get("state", ""),
                items,
                float(o.get("subtotal", 0)),
                float(o.get("shipping", 0)),
                float(o.get("total", 0)),
            ])

        widths = [16, 12, 12, 24, 16, 16, 16, 60, 12, 12, 12]
        for i, w in enumerate(widths, 1):
            ws.column_dimensions[ws.cell(row=1, column=i).column_letter].width = w

        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.border = border
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        buf = io.BytesIO()
        wb.save(buf)
        buf.seek(0)
        return StreamingResponse(
            buf,
            media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            headers={"Content-Disposition": 'attachment; filename="completed_orders.xlsx"'},
        )

    @router.get("/completed/export.pdf")
    async def export_completed_pdf(
        from_date: Optional[str] = Query(None),
        to_date: Optional[str] = Query(None),
        _: dict = Depends(require_admin),
    ):
        orders = (await _completed_orders(from_date, to_date))["orders"]
        pdf = build_pending_orders_pdf(orders)  # same tabular layout
        return StreamingResponse(
            io.BytesIO(pdf),
            media_type="application/pdf",
            headers={"Content-Disposition": 'attachment; filename="completed_orders.pdf"'},
        )

    api.include_router(router)
